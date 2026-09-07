"""Heavy Material Order Request form -> the NR upload CSV.

These arrive as one workbook per order, named for the order reference
("800__DL__M03__01.xlsx"). The form carries a filled "DHL Upload" sheet, and
that sheet is the 47-column order dict nr_csv.transform consumes, column for
column - the same arrangement the Seasonal Treatment Order uses on its "RHPC
Upload" sheet. So this does NOT re-derive anything from the human-facing form:
it reads the sheet the form itself produced, which is the one the requester
checked.

    python heavy_order.py "<800__DL__M03__01.xlsx>" [more...]

WHY THIS NEEDS A REPAIR LAYER

The form computes its own times, and the computation breaks in two ways that
both reach the CSV as rubbish if nothing looks:

1. Excel error literals. collection_time is `=Collection Date + Sheet1!$F$40`,
   and F40 is fed from the Lookup sheet's "Site opening time". Crewe MHD's is
   stored as the TEXT " 08:00" (note the leading space) rather than a time, so
   the addition returns #VALUE!. The closing time is numeric, which is why
   collection_time_end comes through fine and only the start is broken. An
   error literal must never reach the upload: Network Rail's importer splits
   naively on commas and would take "#VALUE!" as a delivery time.

2. Unformatted serials. delivery_time carries number format "General", so the
   cell reads back as 46280.3125 rather than a datetime. The VALUE is correct
   (46280.3125 is 15/09/2026 07:30); only the formatting is missing. Dropping
   it would throw away a good time, so it is converted rather than blanked.

Where a time is genuinely lost, it is rebuilt from the form's OWN Lookup sheet
(site opening / closing hours for that collection site) and reported as a
repair. That is not inventing a value: it is finishing the calculation the
form's formula was already doing, from the form's own data. Anything that
cannot be rebuilt is left blank and shouted about rather than guessed.

Sends nothing. Writes the CSV into the outbox, the same as every other route.
"""
import os
import sys
import datetime

import nr_csv
import outbox

UPLOAD_SHEET = "DHL Upload"
LOOKUP_SHEET = "Lookup"

# The form spells these with underscores; nr_csv reads them with spaces.
TIME_COLS = {
    "collection_time": "collection time",
    "collection_time_end": "collection time end",
    "delivery_time": "delivery time",
    "delivery_time_end": "delivery time end",
}

# Every Excel error literal. None of these may reach the CSV.
ERRORS = ("#VALUE!", "#REF!", "#N/A", "#NAME?", "#DIV/0!", "#NULL!", "#NUM!",
          "#SPILL!", "#CALC!", "#GETTING_DATA")

EPOCH = datetime.datetime(1899, 12, 30)      # Excel's 1900 date system


def is_error(v):
    return isinstance(v, str) and v.strip().upper() in ERRORS


def from_serial(v):
    """An Excel serial -> datetime, or None if it is not one.

    Guarded on a sane range: below 1000 it is far more likely to be a quantity
    than a date, and treating a qty of 150 as 29/05/1900 would be worse than
    leaving it alone.
    """
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        return None
    if not (1000 < float(v) < 80000):
        return None
    return EPOCH + datetime.timedelta(days=float(v))


def _hhmm(v):
    """" 08:00" (however it is stored) -> timedelta, or None."""
    if isinstance(v, datetime.time):
        return datetime.timedelta(hours=v.hour, minutes=v.minute)
    if isinstance(v, datetime.datetime):
        return datetime.timedelta(hours=v.hour, minutes=v.minute)
    if isinstance(v, (int, float)) and 0 <= float(v) < 1:
        return datetime.timedelta(days=float(v))
    s = str(v or "").strip()
    if ":" in s:
        try:
            h, m = s.split(":")[:2]
            return datetime.timedelta(hours=int(h), minutes=int(m[:2]))
        except Exception:
            return None
    return None


def read_lookup(wb):
    """site name -> (opening, closing) as timedeltas, from the form's Lookup."""
    if LOOKUP_SHEET not in wb.sheetnames:
        return {}
    ws = wb[LOOKUP_SHEET]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {}
    hdr = [str(c).strip().lower() if c is not None else "" for c in rows[0]]

    def col(*names):
        for n in names:
            if n in hdr:
                return hdr.index(n)
        return None

    i_name = col("location", " location")
    i_open = col("site opening time")
    i_close = col("site closing time")
    if i_name is None:
        # the header cells carry trailing spaces in the wild - match loosely
        for j, h in enumerate(hdr):
            if h.strip() == "location":
                i_name = j
    out = {}
    for r in rows[1:]:
        if i_name is None or i_name >= len(r) or not r[i_name]:
            continue
        key = str(r[i_name]).strip().lower()
        op = _hhmm(r[i_open]) if i_open is not None and i_open < len(r) else None
        cl = _hhmm(r[i_close]) if i_close is not None and i_close < len(r) else None
        out[key] = (op, cl)
    return out


def _fmt(v):
    if isinstance(v, datetime.datetime):
        return v.strftime("%d/%m/%Y %H:%M")
    if isinstance(v, datetime.date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, datetime.time):
        return v.strftime("%H:%M")
    return "" if v is None else str(v).strip()


def read_orders(path):
    """The DHL Upload sheet's rows as order dicts, plus a repair log."""
    import openpyxl
    import warnings
    warnings.filterwarnings("ignore")
    wb = openpyxl.load_workbook(path, data_only=True)
    if UPLOAD_SHEET not in wb.sheetnames:
        wb.close()
        raise SystemExit(
            f"{os.path.basename(path)}: no '{UPLOAD_SHEET}' sheet - this does not "
            f"look like a Heavy Material Order Request.\n  sheets found: "
            f"{', '.join(wb.sheetnames)}")
    hours = read_lookup(wb)
    ws = wb[UPLOAD_SHEET]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return [], []
    hdr = [str(c).strip() if c is not None else "" for c in rows[0]]

    out, repairs = [], []
    for n, r in enumerate(rows[1:], start=2):
        if not any(c is not None and str(c).strip() for c in r):
            continue
        o, raw = {}, {}
        for i, h in enumerate(hdr):
            if not h:
                continue
            v = r[i] if i < len(r) else None
            key = TIME_COLS.get(h, h)
            timeish = "time" in h.lower() or "date" in h.lower()
            if is_error(v):
                repairs.append((n, h, f"{v} (Excel error)", "blanked"))
                v = None
            elif timeish:
                got = from_serial(v)
                if got is not None:
                    repairs.append((n, h, repr(v), f"read as {got:%d/%m/%Y %H:%M}"))
                    v = got
            o[key] = _fmt(v) if isinstance(
                v, (datetime.datetime, datetime.date, datetime.time)) else v
            raw[h] = v
        o["_raw"] = raw
        if not str(o.get("Customer Order No") or "").strip():
            continue

        # Rebuild a lost collection/delivery time from the form's own Lookup
        # hours. The formula was date + site hours; if the hours cell was text
        # the sum failed, but the date and the hours are both still here.
        site = str(o.get("Site Name - Collection") or "").strip().lower()
        op, cl = hours.get(site, (None, None))
        day = o.get("_raw", {}).get("Collection Date")
        if isinstance(day, datetime.datetime):
            for field, delta, what in (("collection time", op, "site opening"),
                                       ("collection time end", cl, "site closing")):
                if not str(o.get(field) or "").strip() and delta is not None:
                    fixed = day + delta
                    o[field] = fixed.strftime("%d/%m/%Y %H:%M")
                    repairs.append((n, field,
                                    "(lost)", f"rebuilt from {what} -> {o[field]}"))
        out.append(o)
    return out, repairs


def ref_of(orders):
    """The order reference this run is about, for filenames and query subjects."""
    return str(orders[0].get("Customer Order No") or "order").strip() if orders else "order"


def main():
    paths = [a for a in sys.argv[1:] if a.strip() and os.path.exists(a)]
    if not paths:
        print(__doc__)
        print("HEAVY_RESULT orders=0")
        return 2

    stamp = datetime.datetime.now().strftime("%d%m%Y%H%M%S")
    orders, repairs, missing = [], [], []
    for p in paths:
        got, rep = read_orders(p)
        for o in got:
            o["_source"] = os.path.basename(p)
        orders.extend(got)
        repairs.extend((os.path.basename(p),) + r for r in rep)

    if not orders:
        print("  !! no order rows on the upload sheet")
        print("HEAVY_RESULT orders=0")
        return 1

    print(f"Heavy Material Orders - {len(orders)} order(s) from {len(paths)} file(s)\n")
    for o in orders:
        g = lambda k: str(o.get(k) or "").strip()          # noqa: E731
        print(f"  {g('Customer Order No'):22} {g('Site Name - Collection')[:18]:20}"
              f" {g('Postcode'):9} -> {g('Delivery Point')[:16]:18} {g('D Postcode'):9}"
              f" {g('Product Qty')}x {g('Product / Service Code')}")
        print(f"      collection {g('collection time') or '(none)'}"
              f" to {g('collection time end') or '(none)'}")
        print(f"      delivery   {g('delivery time') or '(none)'}"
              f" to {g('delivery time end') or '(none)'}")
        for f in ("collection time", "delivery time"):
            if not g(f):
                missing.append((g("Customer Order No"), f))

    if repairs:
        print(f"\n  {len(repairs)} repair(s) applied to the form's own values:")
        for src, row, field, was, did in repairs:
            print(f"    {src} row {row}: {field} = {was}  ->  {did}")

    # A delivery timed before its own collection cannot be booked, and it is
    # not ours to correct - the requester has to say which end is wrong. Every
    # upload route goes through date_query for this; the email that asks is
    # staged for Review & send and nothing goes on its own.
    raised = ";".join(sorted({str(o.get("Raised by") or "").split(";")[0].strip()
                              for o in orders if o.get("Raised by")}))
    try:
        import date_query
        back = date_query.raise_query(orders, raised,
                                      f"Heavy Material Order {ref_of(orders)}")
    except Exception as ex:
        back = []
        print(f"  !! backwards-date check did not run: {ex}")
    if back:
        print(f"\n  !! {len(back)} DELIVERY BEFORE COLLECTION - not bookable as "
              f"{'they' if len(back) > 1 else 'it'} stands:")
        for p in back:
            print(f"       {p['ref']}  collect {p['collection']}  "
                  f"deliver {p['delivery']}  ({p['back_by']} earlier)")
        print("     An email asking them to confirm is ready for Review & send.")
    print(f"  BACKWARDS_DATES {len(back)}")

    records = nr_csv.transform(orders)
    ref = ref_of(orders).replace("/", "-")
    ref = "".join(ch for ch in ref if ch.isalnum() or ch in "-_")
    csv_out = outbox.path(f"NR_heavy_{ref}_{stamp}.csv")
    nr_csv.write_csv(records, csv_out)
    print(f"\n  CSV: {csv_out}")

    if missing:
        print(f"\n  !! {len(missing)} time(s) could NOT be rebuilt and are BLANK "
              f"on the CSV - fill them in before uploading:")
        for ref_, f in missing:
            print(f"       {ref_}  {f}")

    try:
        import metrics
        metrics.log("heavy_order_built",
                    orders=[str(o.get("Customer Order No")) for o in orders],
                    what=f"{len(repairs)} repair(s)")
    except Exception:
        pass
    print(f"\nHEAVY_RESULT orders={len(orders)} repairs={len(repairs)} "
          f"missing={len(missing)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
