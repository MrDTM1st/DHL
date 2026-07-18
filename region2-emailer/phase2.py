"""
Phase 2 - the reply side of the Region 2 emailer.

For every order the tool has emailed (tracker status 'sent'):
  * find the customer's reply in the DHL inbox
  * if it's an out-of-office / auto-reply -> FLAG it (your cue to chase someone
    else), do NOT count it as a real reply
  * if it's a genuine reply -> parse the delivery details out of it
    (delivery_details) onto the tracker record, AND draft the send-off brief
    into Region 2 > Send Out - a haulier-ready email with the answers in bold
    and the signature attached.

Plus 2-business-day chasers (weekends + England/Wales bank holidays skipped).
Sending chasers is opt-in: preview by default; only sends when told to.

    python phase2.py check          # replies + OOO flags + send-off drafts (no customer email sent)
    python phase2.py chase          # preview who is due a chase (sends nothing)
    python phase2.py chase send     # actually send the chasers
"""
import os, re, sys
import html as _html
from datetime import datetime, date, timedelta
import win32com.client
import build_drafts as bd
import delivery_details as dd
import order_index
import send_order
import tracker

DHL_SMTP = "delali.opoku@dhl.com"
CHASE_AFTER_BDAYS = 2
MAX_CHASES = 3

# England & Wales bank holidays (2026-2027). Extend as needed.
BANK_HOLIDAYS = {
    "2026-01-01", "2026-04-03", "2026-04-06", "2026-05-04", "2026-05-25",
    "2026-08-31", "2026-12-25", "2026-12-28",
    "2027-01-01", "2027-03-26", "2027-03-29", "2027-05-03", "2027-05-31",
    "2027-08-30", "2027-12-27", "2027-12-28",
}


# ---------- dates ----------
def is_working_day(d):
    return d.weekday() < 5 and d.strftime("%Y-%m-%d") not in BANK_HOLIDAYS


def _parse_dt(s):
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(s), fmt)
        except Exception:
            pass
    return None


def business_days_since(dstr):
    dt = _parse_dt(dstr)
    if not dt:
        return 0
    cur, today, n = dt.date(), date.today(), 0
    while cur < today:
        cur += timedelta(days=1)
        if is_working_day(cur):
            n += 1
    return n


# ---------- outlook helpers ----------
def _dhl(ns):
    for i in range(1, ns.Folders.Count + 1):
        if ns.Folders.Item(i).Name.lower() == DHL_SMTP:
            return ns.Folders.Item(i)


def _sub(f, name):
    if f is None:
        return None
    for i in range(1, f.Folders.Count + 1):
        c = f.Folders.Item(i)
        if c.Name.strip().lower() == name.strip().lower():
            return c


def _sendout_folder(ns, create=True):
    dhl = _dhl(ns)
    region2 = _sub(_sub(_sub(dhl, "Inbox"), "Regions"), "Region 2")
    if region2 is None:
        return None
    folder = _sub(region2, "Send out") or _sub(region2, "Send Out")
    if folder is None and create:
        try:
            folder = region2.Folders.Add("Send Out")
        except Exception:
            folder = None
    return folder


def _rt_naive(rt):
    try:
        return datetime.fromtimestamp(rt.timestamp())
    except Exception:
        return None


AUTO_SUBJECT = ("out of office", "automatic reply", "auto-reply", "autoreply",
                "auto reply", "on annual leave", "annual leave", "on holiday",
                "away from the office", "on leave", "out of the office")


def is_auto_reply(item):
    subj = str(getattr(item, "Subject", "") or "").lower()
    if any(k in subj for k in AUTO_SUBJECT):
        return True
    try:
        hdr = item.PropertyAccessor.GetProperty(
            "http://schemas.microsoft.com/mapi/proptag/0x007D001F") or ""
        h = hdr.lower()
        if ("auto-submitted: auto-replied" in h or "auto-submitted:auto-replied" in h
                or "x-autoreply" in h or "x-autorespond" in h
                or "precedence: auto_reply" in h):
            return True
    except Exception:
        pass
    return False


def find_reply(ns, record):
    """The customer's reply for this order: an inbox item received after we
    emailed, whose subject carries one of the order numbers."""
    inbox = _sub(_dhl(ns), "Inbox")
    if inbox is None:
        return None
    emailed = _parse_dt(record.get("emailed_at"))
    to = (record.get("to") or "").strip().lower()
    for order in record.get("orders", []):
        digits = re.sub(r"\D", "", str(order))
        if not digits:
            continue
        try:
            flt = ("@SQL=(\"urn:schemas:httpmail:subject\" LIKE '%" + digits + "%')")
            found = inbox.Items.Restrict(flt)
            found.Sort("[ReceivedTime]", True)
        except Exception:
            found = inbox.Items
        for it in found:
            try:
                rt = _rt_naive(it.ReceivedTime)
                if emailed and rt and rt <= emailed:
                    continue
                sender = str(getattr(it, "SenderEmailAddress", "") or "").lower()
                subj_digits = re.sub(r"\D", "", str(it.Subject or ""))
                if (to and to in sender) or (digits in subj_digits):
                    return it
            except Exception:
                continue
    return None


# ---------- send-off brief ----------
def _load_full(path):
    import openpyxl, warnings
    warnings.filterwarnings("ignore")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h).strip().lower() if h is not None else "" for h in rows[0]]

    def ci(*names):
        for nm in names:
            if nm in hdr:
                return hdr.index(nm)
        return None
    C = dict(order=ci("customer order no"),
             caddr=ci("address1", "address 1"), cpc=ci("postcode"),
             cdate=ci("collection date"),
             daddr=ci("d address1", "d address 1"), dpc=ci("d postcode"),
             ddate=ci("delivery date"),
             prod=ci("product / service code"), prodc=ci("product / description"),
             qty=ci("product qty"))
    return rows[1:], C


def _answer(body, *fragments):
    """Pull the customer's answer that follows one of our questions on the same
    line, e.g. 'Do we need to bring our own offloading? HIAB required' -> the
    text after the '?'. Reading the answer (not the whole body) avoids matching
    keywords that live in the question itself."""
    for line in str(body or "").splitlines():
        low = line.lower()
        if any(fr in low for fr in fragments) and "?" in line:
            ans = line.split("?", 1)[1]
            ans = re.sub(r"^\s*\([^)]*\)\s*", "", ans).strip(" \t-:>")  # drop a trailing bit of the question
            if ans:
                return ans
    return None


def repair_materials():
    """Records built from a BS extract before the column fix stored the NUMERIC
    code as their materials ("6x 0057/063740/0009") instead of the wording. The
    readable description is already sitting in product_codes, so swap it in."""
    d = tracker.load()
    fixed = 0
    for r in d["records"]:
        mats = str(r.get("materials") or "")
        m = re.match(r"^\s*(\d+)\s*x\s*(.+)$", mats)
        if not m or bd._has_words(m.group(2)):
            continue                       # already readable - leave it alone
        words = next((c for c in (r.get("product_codes") or []) if bd._has_words(c)), None)
        if not words:
            continue
        r["materials"] = f"{m.group(1)}x {bd.clean(words)}"
        fixed += 1
    if fixed:
        tracker.save(d)
    return fixed


_ORDER_RE = re.compile(r"\b([5-7]\d{6})\b")
# phrases that mark a genuine delivery-arrangement email TO THE CONTACT
_OUTREACH_MARKERS = (
    "i can get the delivery arranged for you",
    "could you provide me with some details",
    "date and time of delivery?",
    "who will be the contact for",
)
# ...and phrases that mark a haulier request, which must never be enrolled
_HAULIER_MARKERS = ("would you be able to cover", "are you able to cover")


def enrol_untracked(ns, limit=600):
    """SAFETY NET: any order we've emailed the contact about that ISN'T on the
    tracker gets put back on it.

    An order can fall off for several reasons - the wait-list release skips
    enrolment when it finds the order already in Sent Items, an order emailed by
    hand is only enrolled if it happens to be in the CURRENT extract, and a
    record can simply go missing. When that happens the reply is never parsed
    and nothing is ever chased, which breaks the rule that an order is never
    forgotten. This sweeps Sent Items and re-enrols anything missing.

    Skips supplier rails, collection requests, haulier/colleague emails, orders
    already booked, and anything whose delivery date has passed.

    Rebuilding an order from its extract is SLOW, so this runs at most once a
    day and resolves a handful at a time - check() must stay quick."""
    stamp = os.path.join(bd.HERE, "_last_recover.txt")
    try:
        if datetime.fromisoformat(open(stamp, encoding="utf-8").read().strip()) > \
                datetime.now() - timedelta(hours=20):
            return 0
    except Exception:
        pass
    d = tracker.load()
    tracked = {o for r in d["records"] for o in r.get("orders", [])}
    dhl = bd.dhl_store(ns)
    folder = bd.sub(dhl, "Sent Items")
    if folder is None:
        return 0
    items = folder.Items
    try:
        items.Sort("[SentOn]", True)
    except Exception:
        pass
    seen, n = {}, 0
    for it in items:
        n += 1
        if n > limit:
            break
        try:
            subj = str(it.Subject or "")
            bare = re.sub(r"^\s*(re|fw|fwd)\s*:\s*", "", subj, flags=re.I)
            if bare.lower().startswith("collection "):
                continue                      # supplier collection request, not a delivery
            body = str(getattr(it, "Body", "") or "")[:4000].lower()
            # it must be a DELIVERY-ARRANGEMENT email to the contact. You also
            # email hauliers and colleagues about the same order - enrolling
            # those would chase the wrong people entirely.
            if not any(m in body for m in _OUTREACH_MARKERS):
                continue
            if any(m in body for m in _HAULIER_MARKERS):
                continue                      # "would you be able to cover the job below"
            to = str(getattr(it, "To", "") or "")
            if "@dhl.com" in to.lower() or "dhl supply chain" in to.lower():
                continue                      # a colleague, never a delivery contact
            for o in _ORDER_RE.findall(bare):     # only orders named in the SUBJECT
                if o in tracked or o in seen or bd.is_supplier_rail(o):
                    continue
                try:
                    seen[o] = (it.SentOn.strftime("%d/%m/%Y %H:%M"),
                               str(getattr(it, "EntryID", "") or ""))
                except Exception:
                    continue
        except Exception:
            continue
    if not seen:
        return 0
    try:
        open(stamp, "w", encoding="utf-8").write(datetime.now().isoformat())
    except Exception:
        pass
    booked = bd.find_already_emailed(ns, set(seen), limit=1500)
    live = sorted(o for o in seen if not booked.get(o, {}).get("booked"))[:8]
    if not live:
        return 0
    # Rebuild each order from its extract - the same path a chase uses. If it
    # can't be resolved we don't know the date/site/contact, so we don't enrol
    # it (a record we can't chase is worse than no record).
    try:
        collected, _tokens, _nf = send_order.resolve_orders(ns, " ".join(sorted(live)))
        emails = send_order.build_from_collected(collected) if collected else []
    except Exception:
        return 0
    added = 0
    for e in emails:
        ords = [str(o) for o in e.get("orders", [])]
        if not ords or any(o in tracked for o in ords):
            continue
        if not bd._is_future(e.get("date", "")):
            continue                          # delivery date already gone
        when, eid = next((seen[o] for o in ords if o in seen), (None, None))
        tracker.log(orders=ords, to=e.get("to", ""), name=e.get("name", ""),
                    product_codes=e.get("product_codes", []), materials=e.get("materials", ""),
                    site=e.get("site", ""), postcode=e.get("postcode", ""),
                    delivery_date=e.get("date", ""), source="sent (recovered)", status="sent",
                    emailed_at=bd._to_tracker_dt(when) if when else None,
                    only_if_new=True, kind="delivery", orig_entryid=eid)
        tracked.update(ords)
        added += 1
    return added


def product_type_of(record):
    """rails / ballast / sleepers - drives the Moffett lean and the PTS chase."""
    blob = " ".join([record.get("materials", "")] + list(record.get("product_codes") or []))
    return bd.product_type(blob)


_OFFLOAD_WORDS = {"HIAB": "HIAB", "MOFFETT": "Moffett", "SITE/NONE": "Site offloads - none needed",
                  "BOTH": "HIAB or Moffett"}


def brief_lines(record, det, coll="", deliv="", cdate="", mats=""):
    """The haulier brief as (label, value) pairs, built from the PARSED reply.
    A field we don't have is simply left out - no [CHECK] placeholders."""
    date_v = (det.get("date") or {}).get("value") or ""
    lo = (det.get("time") or {}).get("earliest") or ""
    hi = (det.get("time") or {}).get("latest") or ""
    when = " ".join(x for x in (date_v, f"{lo} - {hi}" if lo and hi else lo) if x).strip()

    acc = []
    a = (det.get("artic_access") or {}).get("value")
    veh = (det.get("vehicle") or {}).get("value")
    if a == "yes":
        acc.append("artics can access")
    elif a == "no":
        acc.append("no artics" + (f" - {veh} required" if veh else " - smaller vehicle required"))
    if (det.get("rear_steer") or {}).get("value") == "yes":
        acc.append("rear steer required")

    c = det.get("contact") or {}
    contact = " ".join(x for x in (c.get("name"), c.get("phone")) if x)
    pts = (det.get("pts") or {}).get("value")
    return [
        ("Order", " / ".join(record.get("orders", []))),
        ("Collection", coll),
        ("Delivery", deliv),
        ("Collection date/time", cdate),
        ("Delivery date/time", when),
        ("Materials", mats or record.get("materials", "")),
        ("Vehicle", veh or "(leave with me)"),
        ("Offloading", _OFFLOAD_WORDS.get((det.get("offloading") or {}).get("value"), "")),
        ("Site access", ", ".join(acc)),
        ("Site contact", contact),
        ("What3Words", (det.get("what3words") or {}).get("value") or ""),
        ("PTS", {"yes": "Yes - required", "no": "Not required"}.get(pts, "")),
        ("Notes", (det.get("notes") or {}).get("value") or ""),
    ]


def brief_html(lines):
    """Haulier-ready: the answers in bold so they're easy to scan, then the
    signature. Blank fields are dropped rather than printed empty."""
    body = "".join(f"{_html.escape(k)}: <b>{_html.escape(str(v))}</b><br>"
                   for k, v in lines if str(v).strip())
    return ('<div style="font-family:Calibri,Arial,sans-serif;font-size:11pt;color:#1f1f1f;">'
            "Hi,<br><br>Would you be able to cover the job below;<br><br>"
            f"{body}<br>{bd.SIGNATURE_HTML}</div>")


def build_brief(ns, record, reply_item):
    order = record["orders"][0]
    tmp = os.path.join(bd.HERE, "_brief.xlsx")
    path, fn = order_index.lookup(ns, order, tmp)
    if not path:
        path, fn = send_order.find_extract(ns, order)
    coll = deliv = cdate = ""
    mats = record.get("materials", "")
    if path:
        try:
            rows, C = _load_full(path)
            target = order.split("-")[0]
            grp = [r for r in rows if r[C["order"]] and bd.base_order(r[C["order"]]) == target]
            if grp:
                r0 = grp[0]
                coll = f"{bd.clean(r0[C['caddr']])}, {bd.clean(r0[C['cpc']])}" if C['caddr'] is not None else ""
                deliv = f"{bd.clean(r0[C['daddr']])}, {bd.clean(r0[C['dpc']])}" if C['daddr'] is not None else ""
                cdate = bd.fdate(r0[C['cdate']]) if C['cdate'] is not None else ""
        except Exception:
            pass
    if not deliv:
        deliv = f"{record.get('site', '')}, {record.get('postcode', '')}"

    try:
        reply_body = str(reply_item.Body or "")
    except Exception:
        reply_body = ""
    det = record.get("details") or dd.parse_reply(reply_body, product_type=product_type_of(record))
    lines = brief_lines(record, det, coll, deliv, cdate, mats)

    outlook = win32com.client.Dispatch("Outlook.Application")
    m = outlook.CreateItem(0)
    m.Subject = "SEND OUT: " + " / ".join(record["orders"]) + " " + record.get("site", "")
    bd._attach_qr(m)                      # inline QR used by the signature
    m.HTMLBody = brief_html(lines)        # answers in bold + your signature
    acct = send_order.dhl_account(ns)
    if acct is not None:
        send_order.bind_account(m, acct)
    m.Save()
    folder = _sendout_folder(ns)
    if folder is not None:
        m.Move(folder)
    return True


# ---------- main passes ----------
def _collection_terms():
    """Recipient markers for the collect-first suppliers, from config."""
    terms = set()
    for s in bd.CFG.get("special_collection_suppliers", {}).values():
        for k in s.get("match", []):
            terms.add(str(k).lower())
        for k in s.get("to", []) + s.get("cc", []):
            terms.add(str(k).lower())
    return terms


def enrol_collection(ns):
    """Auto-track the supplier collection-request emails you've sent (subject
    'Collection <order> ...' TO Anderton / BCM / Trough Tec) so they get chased if
    the supplier doesn't come back with the collection details. Gated on a
    special-collection supplier recipient AND a 7-digit order number, so ordinary
    mail never lands on the tracker. Added once (only_if_new); chased in-thread.
    Returns how many were newly enrolled."""
    dhl = _dhl(ns)
    sent = _sub(dhl, "Sent Items")
    terms = _collection_terms()
    if sent is None or not terms:
        return 0
    items = sent.Items
    try:
        items.Sort("[SentOn]", True)
    except Exception:
        pass
    existing = {r["id"] for r in tracker.load().get("records", [])}
    added = n = 0
    for it in items:
        n += 1
        if n > 400:
            break
        try:
            core = re.sub(r"^\s*(re|fw|fwd)\s*:\s*", "",
                          str(getattr(it, "Subject", "") or "").strip(), flags=re.I)
            if not core.lower().startswith("collection "):
                continue
            who = (str(getattr(it, "To", "") or "") + " " + str(getattr(it, "CC", "") or "")).lower()
            if not any(t in who for t in terms):
                continue                       # only the collect-first suppliers
            m = re.search(r"(?<!\d)(\d{7})(?!\d)", core)
            if not m:
                continue                       # need a real order number to track/chase
            order = m.group(1)
            rid = tracker._key([order], "")
            if rid in existing:
                continue
            when = None
            try:
                s = it.SentOn
                when = datetime(s.year, s.month, s.day, s.hour, s.minute).strftime("%Y-%m-%d %H:%M")
            except Exception:
                pass
            tracker.log(orders=[order], to=str(getattr(it, "To", "") or ""), name="",
                        product_codes=[], materials="collection details", site="", postcode="",
                        delivery_date="", source="collection", status="sent",
                        emailed_at=when, only_if_new=True, kind="collection",
                        orig_entryid=str(getattr(it, "EntryID", "") or ""))
            existing.add(rid)
            added += 1
        except Exception:
            continue
    return added


def check(ns=None):
    ns = ns or bd.get_ns()
    enrol_collection(ns)          # track supplier collection emails before scanning for replies
    repaired = repair_materials()     # BS-file rows that stored the numeric code (fast, no COM)
    d = tracker.load()
    replies = ooo = briefs = 0
    for r in d["records"]:
        if r.get("status") != "sent":
            continue
        if r.get("reply_at") or r.get("ooo_at"):
            continue
        item = find_reply(ns, r)
        if not item:
            continue
        if is_auto_reply(item):
            r["ooo_at"] = tracker._now()
            ooo += 1
        else:
            r["reply_at"] = tracker._now()
            replies += 1
            # parse their answers into the structured fields CTMS needs, so the
            # brief is pre-filled and the chaser can ask for ONLY what's missing
            try:
                pt = product_type_of(r)
                r["details"] = dd.parse_reply(str(item.Body or ""), product_type=pt)
                r["missing"] = dd.missing(r["details"], pt)
            except Exception as e:
                r["details_note"] = str(e)[:140]
            if r.get("kind") == "collection":
                continue   # supplier replied with collection details - no delivery send-off brief
            try:
                if build_brief(ns, r, item):
                    r["sendoff_ready"] = True
                    briefs += 1
            except Exception as e:
                r["sendoff_note"] = str(e)[:140]
    # remove orders YOU'VE booked yourself - your "this order has been arranged
    # with ..." email (incl. a MAN ref) in Sent Items - so they stop being tracked
    # and never get chased. This is the "if you say an order's booked, drop it"
    # rule. Scan deep (1500): tracked orders can be weeks old, so their booking
    # email sits well down the Sent Items list.
    all_orders = {o for r in d["records"] for o in r.get("orders", [])}
    booked = bd.find_already_emailed(ns, all_orders, limit=1500) if all_orders else {}
    before = len(d["records"])
    d["records"] = [r for r in d["records"]
                    if not (r.get("orders") and any(booked.get(o, {}).get("booked") for o in r["orders"]))]
    booked_removed = before - len(d["records"])
    removed = tracker.drop_completed(d)   # completed orders leave the tracker
    tracker.save(d)
    print(f"check: {replies} new repl(y/ies), {ooo} out-of-office flagged, "
          f"{briefs} send-off draft(s) created, {booked_removed} booked-by-you removed, "
          f"{removed} completed order(s) removed, {repaired} product wording(s) repaired.")
    return replies, ooo, briefs


def _due_for_chase(r):
    if r.get("status") != "sent":
        return False
    if r.get("ooo_at"):
        return False
    # a reply only closes it if it actually answered everything - a PARTIAL
    # reply still gets chased, but only for the fields it left blank
    if r.get("reply_at") and not r.get("missing"):
        return False
    if r.get("chases", 0) >= MAX_CHASES:
        return False
    if business_days_since(r.get("emailed_at")) < CHASE_AFTER_BDAYS:
        return False
    if r.get("last_chased_at") == date.today().isoformat():
        return False
    return True


def _chase_in_thread(ns, record):
    """Follow up on the exact email you sent, kept on the same thread. Used for
    supplier collection requests - there is no extract to rebuild from, so we
    reply to the original Sent item (quoting it) asking for the details."""
    eid = record.get("orig_entryid")
    if not eid:
        return False
    try:
        item = ns.GetItemFromID(eid)
    except Exception:
        return False
    if item is None:
        return False
    try:
        reply = item.Reply()
        note = ("Hi,\n\nJust following up on the below - could I please get the "
                "collection details when you have a moment?\n\n")
        try:
            reply.Body = note + reply.Body
        except Exception:
            pass
        acct = send_order.dhl_account(ns)
        if acct is not None:
            send_order.bind_account(reply, acct)
        reply.Send()
        return True
    except Exception:
        return False


def _send_chase(ns, record):
    if record.get("kind") == "collection":
        return _chase_in_thread(ns, record)
    collected, tokens, nf = send_order.resolve_orders(ns, " ".join(record["orders"]))
    if not collected:
        return False
    emails = send_order.build_from_collected(collected)
    if not emails:
        return False
    e = emails[0]
    original = e["message"].split("\n\n", 1)[-1]
    greet = f"Hi {e['name']}," if e.get("name") else "Hi,"
    # if they replied but left gaps, ask for EXACTLY those - not "any update?"
    miss = record.get("missing") or []
    if miss and record.get("reply_at"):
        need = miss[0] if len(miss) == 1 else ", ".join(miss[:-1]) + " and " + miss[-1]
        ask = f"Thanks for coming back to me - I just need the {need} and I can get this booked."
    else:
        ask = "Can I please get a reply to the below?"
    e["message"] = f"{greet}\n\n{ask}\n\n{original}"
    e["html"] = bd.html_from_message(e["message"])
    e["cc"] = ""
    return send_order.send_emails(ns, [e]) > 0   # send_emails -> tracker.log bumps chases


_LOCK_SOCK = None


def _chase_lock():
    """Only ONE chase run at a time. Both agents (local + cloud) fire
    `phase2.py chase send`, and without this they raced and every contact got
    the same chaser twice, a second apart."""
    global _LOCK_SOCK
    import socket
    s = socket.socket()
    try:
        s.bind(("127.0.0.1", 8790))
    except OSError:
        return False
    _LOCK_SOCK = s          # held for the life of the process
    return True


def _claim(rid):
    """Mark this record chased TODAY *before* sending, so a concurrent or
    immediately-following run skips it. Claiming first means a failed send
    costs us a day's delay - far better than emailing someone twice."""
    d = tracker.load()
    today = date.today().isoformat()
    for r in d["records"]:
        if r.get("id") == rid:
            if r.get("last_chased_at") == today:
                return False                 # already claimed by another run
            r["last_chased_at"] = today
            tracker.save(d)
            return True
    return False


def run_chasers(ns=None, send=False):
    if send and not _chase_lock():
        print("chasers: another chase run is already in progress - skipping.")
        return []
    ns = ns or bd.get_ns()
    d = tracker.load()
    due = [r for r in d["records"] if _due_for_chase(r)]
    out = []
    chased_ids = []
    for r in due:
        bd_n = business_days_since(r.get("emailed_at"))
        if not send:
            out.append(f"  DUE  {' / '.join(r['orders'])} -> {r['to']} "
                       f"(bday {bd_n}, would be chase #{r.get('chases', 0) + 1})")
            continue
        if not _claim(r["id"]):
            out.append(f"  SKIP {' / '.join(r['orders'])} -> already chased today")
            continue
        ok = _send_chase(ns, r)
        out.append(f"  {'SENT' if ok else 'FAIL'} {' / '.join(r['orders'])} -> {r['to']}")
        if ok:
            chased_ids.append(r["id"])
    if send and chased_ids:
        d2 = tracker.load()
        for r in d2["records"]:
            if r["id"] in chased_ids and r.get("kind") == "collection":
                r["chases"] = r.get("chases", 0) + 1   # in-thread chase skips tracker.log's bump
        tracker.save(d2)
    verb = "sent" if send else "due"
    print(f"chasers: {len(chased_ids) if send else len(due)} {verb}.")
    print("\n".join(out) if out else "  (none)")
    return out


def learn_detail(rec_id, field, value):
    """Confirm/correct one parsed field. The wording that produced it is
    remembered, so the same phrasing is never guessed again."""
    d = tracker.load()
    for r in d["records"]:
        if r.get("id") != rec_id:
            continue
        det = r.get("details") or {}
        cell = det.get(field)
        if not isinstance(cell, dict):
            return f"no parsed '{field}' on {rec_id}"
        raw = cell.get("raw") or ""
        if raw:
            dd.learn(field, raw, value)          # phrase -> value, permanently
        if field == "contact":
            cell["name"] = value
        elif field == "time":
            cell["earliest"] = value
        else:
            cell["value"] = value
        cell["confidence"] = dd.HIGH
        r["missing"] = dd.missing(det, product_type_of(r))
        tracker.save(d)
        return f"learned {field}={value!r} for {rec_id} (from {raw[:40]!r})"
    return f"record {rec_id} not found"


def main():
    cmd = sys.argv[1] if len(sys.argv) > 1 else "check"
    if cmd == "learn":                     # no Outlook needed - pure tracker edit
        if len(sys.argv) < 5:
            print("usage: phase2.py learn <record id> <field> <value>"); return
        print(learn_detail(sys.argv[2], sys.argv[3], " ".join(sys.argv[4:])))
        return
    ns = bd.get_ns()
    if cmd == "recover":
        # slow (rebuilds each order from its extract) - runs on its own daily
        # cadence, never inside check()
        print(f"recovered {enrol_untracked(ns)} untracked order(s).")
    elif cmd == "check":
        check(ns)
    elif cmd == "chase":
        run_chasers(ns, send=(len(sys.argv) > 2 and sys.argv[2] == "send"))
    elif cmd == "all":
        check(ns)
        run_chasers(ns, send=False)
    else:
        print("usage: phase2.py check | chase [send] | all | learn <id> <field> <value>")


if __name__ == "__main__":
    main()
